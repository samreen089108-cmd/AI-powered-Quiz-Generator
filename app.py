"""
AI-Powered Quiz Generator from Notes
Stack: Python + Streamlit + Ollama (local LLM, RAG) + HTML/CSS
Features: PDF/Word/TXT upload, flashcards (front/back + next), graded quiz,
          chat with notes, download as PDF / Word / HTML / TXT / JSON.
"""
import html
import io
import json
import re

import numpy as np
import ollama
import streamlit as st
from docx import Document
from fpdf import FPDF
from pypdf import PdfReader

st.set_page_config(page_title="AI Quiz Generator", page_icon="📝", layout="wide")

# ------------------------------------------------------------------ CSS
CSS = """
.flashcard {
    border-radius: 20px; padding: 2.2rem 2.4rem; min-height: 280px; color: #fff;
    box-shadow: 0 10px 30px rgba(0,0,0,.25); animation: pop .35s ease;
}
.flashcard.front { background: linear-gradient(135deg, #4f46e5, #7c3aed); }
.flashcard.back  { background: linear-gradient(135deg, #047857, #10b981); }
.flashcard h2 { color: #fff; font-size: 1.6rem; margin: .6rem 0 1rem 0; line-height: 1.4; }
.flashcard ul { list-style: none; padding: 0; margin: 0; }
.flashcard li {
    background: rgba(255,255,255,.16); border-radius: 10px;
    padding: .55rem .9rem; margin-bottom: .5rem; font-size: 1.05rem;
}
.flashcard .tag { font-size: .8rem; letter-spacing: .12em; opacity: .85; font-weight: 700; }
.flashcard .answer { font-size: 1.5rem; font-weight: 700; margin: .6rem 0 1rem 0; }
.flashcard .why { font-size: 1.05rem; line-height: 1.5; }
.flashcard .src {
    margin-top: 1.2rem; font-size: .85rem; opacity: .9; font-style: italic;
    border-top: 1px solid rgba(255,255,255,.35); padding-top: .7rem;
}
.status-ok  { color: #10b981; font-weight: 600; }
.status-bad { color: #ef4444; font-weight: 600; }
@keyframes pop { from { transform: scale(.97); opacity: .4; } to { transform: scale(1); opacity: 1; } }
"""
st.markdown("<style>" + CSS + "</style>", unsafe_allow_html=True)


def esc(x) -> str:
    return html.escape(str(x))


# ------------------------------------------------------------------ Ollama
def installed_models():
    try:
        res = ollama.list()
        models = res["models"] if isinstance(res, dict) else res.models
        return [(m["model"] if isinstance(m, dict) else m.model) for m in models]
    except Exception:
        return []


def chat_llm(model, messages, json_mode=False):
    kwargs = {"model": model, "messages": messages, "options": {"temperature": 0.3}}
    if json_mode:
        kwargs["format"] = "json"
    return ollama.chat(**kwargs)["message"]["content"]


def embed_texts(model, texts, batch=16):
    out = []
    for i in range(0, len(texts), batch):
        out += ollama.embed(model=model, input=texts[i:i + batch])["embeddings"]
    return out


# ------------------------------------------------------------------ reading files
def read_file(f) -> str:
    name = f.name.lower()
    if name.endswith(".pdf"):
        return "\n".join((p.extract_text() or "") for p in PdfReader(f).pages)
    if name.endswith(".docx"):
        doc = Document(f)
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:  # include text inside tables too
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    return f.read().decode("utf-8", errors="ignore")  # txt, md


# ------------------------------------------------------------------ RAG
def chunk_text(text, size=800, overlap=150):
    text = re.sub(r"\s+", " ", text).strip()
    chunks, start = [], 0
    while start < len(text):
        chunks.append(text[start:start + size])
        start += size - overlap
    return chunks


def build_index(texts, embed_model):
    chunks, sources = [], []
    for fname, txt in texts.items():
        for c in chunk_text(txt):
            chunks.append(c)
            sources.append(fname)
    try:
        emb = np.array(embed_texts(embed_model, chunks), dtype=float)
        emb = emb / np.maximum(np.linalg.norm(emb, axis=1, keepdims=True), 1e-9)
    except Exception:
        emb = None  # fall back to keyword search
    st.session_state.update(chunks=chunks, sources=sources, emb=emb, embed_model=embed_model)


def keyword_scores(query):
    q = set(re.findall(r"\w+", query.lower()))
    return [len(q & set(re.findall(r"\w+", c.lower()))) for c in st.session_state.chunks]


def retrieve(query, k=5):
    scores = None
    if st.session_state.emb is not None:
        try:
            qv = np.array(embed_texts(st.session_state.embed_model, [query])[0], dtype=float)
            qv = qv / max(np.linalg.norm(qv), 1e-9)
            scores = (st.session_state.emb @ qv).tolist()
        except Exception:
            scores = None
    if scores is None:
        scores = keyword_scores(query)
    top = np.argsort(scores)[::-1][:k]
    return [(st.session_state.chunks[i], st.session_state.sources[i]) for i in top]


def spread_chunks(n):
    total = len(st.session_state.chunks)
    n = min(n, total)
    idx = sorted({int(round(i)) for i in np.linspace(0, total - 1, n)})
    return [(st.session_state.chunks[i], st.session_state.sources[i]) for i in idx]


# ------------------------------------------------------------------ quiz generation
def strip_prefix(s):
    return re.sub(r"^[A-Da-d][\.\)]\s+", "", str(s).strip())


def clean_questions(data):
    if isinstance(data, dict):
        data = data.get("questions") or next((v for v in data.values() if isinstance(v, list)), [])
    good = []
    for it in data if isinstance(data, list) else []:
        try:
            q = str(it["question"]).strip()
            opts = [strip_prefix(o) for o in it["options"]][:4]
            ans = str(it["answer"]).strip()
            if len(ans) == 1 and ans.upper() in "ABCD" and "ABCD".index(ans.upper()) < len(opts):
                ans = opts["ABCD".index(ans.upper())]
            ans = strip_prefix(ans)
            if ans not in opts:
                match = [o for o in opts if o.lower() == ans.lower()]
                if not match:
                    continue
                ans = match[0]
            if len(opts) < 2 or not q:
                continue
            good.append({"question": q, "options": opts, "answer": ans,
                         "explanation": str(it.get("explanation", "")),
                         "source_snippet": str(it.get("source_snippet", ""))})
        except (KeyError, TypeError):
            continue
    return good


def generate_quiz(model, n, difficulty, topic):
    ctx = retrieve(topic, k=8) if topic.strip() else spread_chunks(10)
    context = "\n\n".join(f"[Source: {s}]\n{c}" for c, s in ctx)
    prompt = f"""Use ONLY the notes below. Do not use outside knowledge.

NOTES:
{context}

Create {n} multiple-choice questions at {difficulty} difficulty.
Reply with JSON only, in exactly this shape:
{{"questions": [{{"question": "...", "options": ["...", "...", "...", "..."],
"answer": "exact text of the correct option", "explanation": "why it is correct",
"source_snippet": "short quote from the notes"}}]}}"""
    for _ in range(2):  # small local models sometimes need a second try
        raw = chat_llm(model, [{"role": "user", "content": prompt}], json_mode=True)
        try:
            quiz = clean_questions(json.loads(raw))
        except json.JSONDecodeError:
            quiz = []
        if quiz:
            return quiz
    raise ValueError("The model did not return valid questions. Try again or use a bigger model.")


# ------------------------------------------------------------------ exports
def pdf_safe(t):
    t = str(t).encode("latin-1", "replace").decode("latin-1")  # built-in PDF fonts are Latin-1
    return re.sub(r"(\S{60})", r"\1 ", t)  # break very long words so they fit


def quiz_to_pdf(quiz, with_answers):
    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Quiz" + (" - Answer Key" if with_answers else ""), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)
    for i, q in enumerate(quiz, 1):
        pdf.set_font("Helvetica", "B", 11)
        pdf.multi_cell(0, 6, pdf_safe(f"Q{i}. {q['question']}"), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 11)
        for j, o in enumerate(q["options"]):
            pdf.multi_cell(0, 6, pdf_safe(f"    {'ABCD'[j]}. {o}"), new_x="LMARGIN", new_y="NEXT")
        if with_answers:
            pdf.set_font("Helvetica", "I", 10)
            pdf.multi_cell(0, 5, pdf_safe(f"Answer: {q['answer']}"), new_x="LMARGIN", new_y="NEXT")
            pdf.multi_cell(0, 5, pdf_safe(f"Why: {q['explanation']}"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
    return bytes(pdf.output())


def quiz_to_docx(quiz, with_answers):
    doc = Document()
    doc.add_heading("Quiz" + (" - Answer Key" if with_answers else ""), 0)
    for i, q in enumerate(quiz, 1):
        p = doc.add_paragraph()
        p.add_run(f"Q{i}. {q['question']}").bold = True
        for j, o in enumerate(q["options"]):
            doc.add_paragraph(f"{'ABCD'[j]}. {o}")
        if with_answers:
            p = doc.add_paragraph()
            p.add_run(f"Answer: {q['answer']}").italic = True
            doc.add_paragraph(f"Why: {q['explanation']}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


EXPORT_CSS = """
body{font-family:Segoe UI,Arial,sans-serif;max-width:800px;margin:2rem auto;color:#111;padding:0 1rem}
h1{color:#4f46e5} .q{margin:1.4rem 0;page-break-inside:avoid}
.q h3{margin:.2rem 0 .5rem 0} ol{list-style:upper-alpha;margin:.3rem 0 .6rem 1.4rem}
.ans{background:#ecfdf5;border-left:4px solid #10b981;padding:.5rem .8rem;border-radius:4px}
@media print{body{margin:0}}
"""


def quiz_to_html(quiz, with_answers):
    parts = ["<!DOCTYPE html><html><head><meta charset='utf-8'><title>Quiz</title><style>",
             EXPORT_CSS, "</style></head><body><h1>Quiz", " - Answer Key" if with_answers else "", "</h1>"]
    for i, q in enumerate(quiz, 1):
        opts = "".join(f"<li>{esc(o)}</li>" for o in q["options"])
        parts.append(f"<div class='q'><h3>Q{i}. {esc(q['question'])}</h3><ol>{opts}</ol>")
        if with_answers:
            parts.append(f"<div class='ans'><b>Answer:</b> {esc(q['answer'])}<br>{esc(q['explanation'])}</div>")
        parts.append("</div>")
    parts.append("</body></html>")
    return "".join(parts)


def quiz_to_text(quiz, with_answers):
    out = []
    for i, q in enumerate(quiz, 1):
        out.append(f"Q{i}. {q['question']}")
        out += [f"   {'ABCD'[j]}. {o}" for j, o in enumerate(q["options"])]
        if with_answers:
            out.append(f"   Answer: {q['answer']}\n   Why: {q['explanation']}")
        out.append("")
    return "\n".join(out)


# ------------------------------------------------------------------ flashcard HTML
def card_html(q, i, total, flipped):
    if not flipped:
        opts = "".join(f"<li><b>{'ABCD'[j]}.</b> {esc(o)}</li>" for j, o in enumerate(q["options"]))
        return (f"<div class='flashcard front'><div class='tag'>QUESTION {i + 1} / {total}</div>"
                f"<h2>{esc(q['question'])}</h2><ul>{opts}</ul></div>")
    src = f"<div class='src'>📄 From your notes: “{esc(q['source_snippet'])}”</div>" if q["source_snippet"] else ""
    return (f"<div class='flashcard back'><div class='tag'>ANSWER</div>"
            f"<div class='answer'>✅ {esc(q['answer'])}</div>"
            f"<div class='why'>{esc(q['explanation'])}</div>{src}</div>")


# ------------------------------------------------------------------ state
for k, v in {"quiz": [], "card": 0, "flipped": False, "submitted": False, "chat": []}.items():
    st.session_state.setdefault(k, v)

# ------------------------------------------------------------------ sidebar
with st.sidebar:
    st.header("⚙️ Setup")
    models = installed_models()
    if models:
        st.markdown("<span class='status-ok'>● Ollama connected</span>", unsafe_allow_html=True)
        gen_model = st.selectbox("Model", models,
                                 index=next((i for i, m in enumerate(models) if m.startswith("llama3.2")), 0))
    else:
        st.markdown("<span class='status-bad'>● Ollama not found</span>", unsafe_allow_html=True)
        st.caption("Start Ollama, then refresh this page.")
        gen_model = st.text_input("Model", "llama3.2")
    embed_model = st.text_input("Embedding model", "nomic-embed-text")

    files = st.file_uploader("Upload notes (PDF, Word, TXT, MD)",
                             type=["pdf", "docx", "txt", "md"], accept_multiple_files=True)
    pasted = st.text_area("…or paste notes", height=110)

    if st.button("📚 Process notes", use_container_width=True):
        texts = {f.name: read_file(f) for f in files} if files else {}
        if pasted.strip():
            texts["pasted_notes"] = pasted
        texts = {k: v for k, v in texts.items() if v.strip()}
        if not texts:
            st.error("Add some notes first (scanned PDFs without text can't be read).")
        else:
            with st.spinner("Reading and indexing your notes..."):
                build_index(texts, embed_model)
            msg = f"Indexed {len(st.session_state.chunks)} chunks from {len(texts)} source(s)."
            if st.session_state.emb is None:
                st.warning(msg + " Embedding model unavailable, using keyword search instead.")
            else:
                st.success(msg)

    st.divider()
    st.subheader("Quiz settings")
    n_q = st.slider("Number of questions", 3, 15, 5)
    difficulty = st.selectbox("Difficulty", ["easy", "medium", "hard"], index=1)
    topic = st.text_input("Focus topic (optional)")
    if st.button("✨ Generate quiz", type="primary", use_container_width=True):
        if "chunks" not in st.session_state:
            st.error("Process your notes first.")
        else:
            with st.spinner("Generating questions locally… this can take a minute."):
                try:
                    st.session_state.quiz = generate_quiz(gen_model, n_q, difficulty, topic)
                    st.session_state.card = 0
                    st.session_state.flipped = False
                    st.session_state.submitted = False
                    for key in [k for k in st.session_state if k.startswith("ans_")]:
                        del st.session_state[key]
                    st.success(f"Created {len(st.session_state.quiz)} questions.")
                except Exception as e:
                    st.error(f"Could not generate quiz: {e}")

# ------------------------------------------------------------------ main
st.title("📝 AI-Powered Quiz Generator from Notes")
st.caption("Runs fully on your laptop with Ollama. Your notes never leave your machine.")
tab_cards, tab_quiz, tab_chat, tab_dl = st.tabs(
    ["🃏 Flashcards", "✅ Take Quiz", "💬 Chat with Notes", "⬇️ Download"])
quiz = st.session_state.quiz

# ---- Flashcards: question on front, answer on back, next button
with tab_cards:
    if not quiz:
        st.info("Upload notes, process them, then generate a quiz from the sidebar.")
    else:
        i = st.session_state.card
        st.progress((i + 1) / len(quiz))
        st.markdown(card_html(quiz[i], i, len(quiz), st.session_state.flipped), unsafe_allow_html=True)
        st.write("")
        c1, c2, c3 = st.columns(3)
        if c1.button("⬅️ Previous", disabled=i == 0, use_container_width=True):
            st.session_state.card -= 1
            st.session_state.flipped = False
            st.rerun()
        label = "🔄 Show question" if st.session_state.flipped else "🔄 Show answer"
        if c2.button(label, use_container_width=True):
            st.session_state.flipped = not st.session_state.flipped
            st.rerun()
        if c3.button("Next ➡️", disabled=i == len(quiz) - 1, use_container_width=True):
            st.session_state.card += 1
            st.session_state.flipped = False
            st.rerun()

# ---- Graded quiz
with tab_quiz:
    if not quiz:
        st.info("Generate a quiz first.")
    else:
        for i, q in enumerate(quiz):
            st.markdown(f"**Q{i + 1}. {q['question']}**")
            st.radio("Choose one", q["options"], index=None, key=f"ans_{i}",
                     disabled=st.session_state.submitted, label_visibility="collapsed")
            if st.session_state.submitted:
                if st.session_state.get(f"ans_{i}") == q["answer"]:
                    st.success("Correct ✅")
                else:
                    st.error(f"Wrong ❌ Correct answer: {q['answer']}")
                st.caption(q["explanation"])
            st.divider()
        if not st.session_state.submitted:
            if st.button("Submit quiz", type="primary"):
                st.session_state.submitted = True
                st.rerun()
        else:
            score = sum(st.session_state.get(f"ans_{i}") == q["answer"] for i, q in enumerate(quiz))
            st.metric("Your score", f"{score} / {len(quiz)}")

# ---- Chat with notes
with tab_chat:
    if "chunks" not in st.session_state:
        st.info("Process your notes first.")
    else:
        for m in st.session_state.chat:
            with st.chat_message(m["role"]):
                st.markdown(m["content"])
        if question := st.chat_input("Ask anything about your notes…"):
            st.session_state.chat.append({"role": "user", "content": question})
            with st.chat_message("user"):
                st.markdown(question)
            ctx = retrieve(question, k=4)
            context = "\n\n".join(f"[{s}] {c}" for c, s in ctx)
            system = {"role": "system", "content":
                      "Answer using ONLY the notes below. If the answer is not in the notes, say so.\n\n"
                      f"NOTES:\n{context}"}
            with st.chat_message("assistant"):
                with st.spinner("Thinking…"):
                    try:
                        answer = chat_llm(gen_model, [system] + st.session_state.chat[-8:])
                    except Exception as e:
                        answer = f"Could not reach Ollama: {e}"
                st.markdown(answer)
                with st.expander("Sources used"):
                    for c, s in ctx:
                        st.caption(f"**{s}**: {c[:300]}…")
            st.session_state.chat.append({"role": "assistant", "content": answer})

# ---- Downloads
with tab_dl:
    if not quiz:
        st.info("Generate a quiz first.")
    else:
        c1, c2 = st.columns(2)
        with c1:
            st.markdown("**Quiz (no answers)**")
            st.download_button("📕 PDF", quiz_to_pdf(quiz, False), "quiz.pdf",
                               "application/pdf", use_container_width=True)
            st.download_button("📘 Word (.docx)", quiz_to_docx(quiz, False), "quiz.docx",
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               use_container_width=True)
            st.download_button("🌐 HTML", quiz_to_html(quiz, False), "quiz.html",
                               "text/html", use_container_width=True)
            st.download_button("📄 TXT", quiz_to_text(quiz, False), "quiz.txt",
                               use_container_width=True)
        with c2:
            st.markdown("**Answer key**")
            st.download_button("🔑 PDF", quiz_to_pdf(quiz, True), "quiz_answers.pdf",
                               "application/pdf", use_container_width=True)
            st.download_button("🔑 Word (.docx)", quiz_to_docx(quiz, True), "quiz_answers.docx",
                               "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               use_container_width=True)
            st.download_button("🔑 HTML", quiz_to_html(quiz, True), "quiz_answers.html",
                               "text/html", use_container_width=True)
            st.download_button("🧾 JSON (full data)", json.dumps(quiz, indent=2), "quiz.json",
                               "application/json", use_container_width=True)
